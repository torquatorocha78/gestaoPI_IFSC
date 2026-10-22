import hashlib
import io
import os
import re
import time
import unicodedata
from datetime import datetime

import requests
import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

import database as db


MODEL = os.getenv("GROQ_MODEL", "openai/gpt-oss-20b")


def _secret(name, default=""):
    value = os.getenv(name)
    if value:
        return value
    try:
        return str(st.secrets.get(name, default))
    except Exception:
        return default


def _sem_acentos(value):
    normalized = unicodedata.normalize("NFKD", value or "")
    return "".join(char for char in normalized if not unicodedata.combining(char)).casefold()


def _excel_column(number):
    """Converte um índice de coluna baseado em 1 para letras do Excel."""
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _column_number(label):
    number = 0
    for char in label.upper():
        number = number * 26 + ord(char) - 64
    return number


def _is_empty_or_not_applicable(value):
    value = re.sub(r"\s+", " ", (value or "")).strip()
    normalized = _sem_acentos(value).strip(" .!;\t")
    return not normalized or normalized in {
        "nao aplicavel",
        "nao se aplica",
        "n/a",
    }


def _read_spreadsheet(uploaded_file):
    """Lê uma planilha Excel/ODS com perguntas na primeira linha e respostas na segunda."""
    data = uploaded_file.getvalue()
    extension = os.path.splitext(uploaded_file.name)[1].casefold()
    engines = {
        ".ods": "odf",
        ".xls": "xlrd",
        ".xlsx": "openpyxl",
        ".xlsm": "openpyxl",
        ".xlsb": "pyxlsb",
        ".xltx": "openpyxl",
        ".xltm": "openpyxl",
    }
    engine = engines.get(extension)
    if not engine:
        raise ValueError("Formato não suportado. Envie XLSX, XLS, XLSM, XLSB, XLTX, XLTM ou ODS.")

    try:
        import pandas as pd
        frame = pd.read_excel(
            io.BytesIO(data),
            engine=engine,
            header=None,
            dtype=str,
            keep_default_na=False,
        )
        rows = frame.values.tolist()
    except ImportError as exc:
        package = {"odf": "odfpy", "xlrd": "xlrd", "openpyxl": "openpyxl", "pyxlsb": "pyxlsb"}[engine]
        raise ValueError(f"Para abrir este formato, instale a dependência {package} do projeto.") from exc
    except Exception as exc:
        raise ValueError(f"Não consegui ler a planilha {extension}: {exc}") from exc

    rows = [[str(cell or "").strip() for cell in row] for row in rows]
    rows = [row for row in rows if any(cell for cell in row)]
    if len(rows) != 2:
        raise ValueError(
            f"Esperava duas linhas lógicas (perguntas e respostas), mas encontrei {len(rows)}. "
            "Exporte uma única resposta do formulário, com perguntas na linha 1 e respostas na linha 2."
        )

    questions, answers = rows
    if len(questions) != len(answers):
        raise ValueError(
            f"A linha de perguntas tem {len(questions)} colunas e a de respostas tem {len(answers)}. "
            "Confira se a planilha contém uma única resposta completa."
        )

    pairs = []
    for index, (question, answer) in enumerate(zip(questions, answers), 1):
        question = re.sub(r"\s+", " ", (question or "")).strip()
        answer = (answer or "").strip()
        if not question or _is_empty_or_not_applicable(answer):
            continue
        pairs.append({
            "coluna": _excel_column(index),
            "pergunta": question,
            "resposta": answer,
        })
    if not pairs:
        raise ValueError("A planilha não tem respostas preenchidas diferentes de 'Não aplicável'.")
    return pairs


def _is_patent_form(pairs):
    for item in pairs:
        question = _sem_acentos(item["pergunta"])
        answer = _sem_acentos(item["resposta"])
        if "invencao" not in question:
            continue
        if "programa de computador" in answer or "software" in answer:
            return False
        if "desenho industrial" in answer:
            return False
        if "patente" in answer or "invencao" in answer or "modelo de utilidade" in answer:
            return True
    return None


def _filter_patent_pairs(pairs):
    """Retém dados de patente e remove software, desenho e metadados do formulário."""
    result = []
    excluded_start = _column_number("EI")
    excluded_end = _column_number("UO")
    excluded_metadata_columns = {2, 6}  # Data de envio e data da última ação.

    for item in pairs:
        column_number = _column_number(item["coluna"])
        if column_number in excluded_metadata_columns:
            continue
        if excluded_start <= column_number <= excluded_end:
            continue
        result.append(item)
    return result


def _format_pairs(pairs):
    return "\n\n".join(
        f"[{item['coluna']}] PERGUNTA: {item['pergunta']}\nRESPOSTA: {item['resposta']}"
        for item in pairs
    )


def _groq_request(api_key, messages, max_tokens):
    request_body = {
        "model": MODEL,
        "temperature": 0.1,
        "max_tokens": max_tokens,
        "messages": messages,
    }
    headers = {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}

    for attempt in range(3):
        response = requests.post(
            "https://api.groq.com/openai/v1/chat/completions",
            headers=headers,
            json=request_body,
            timeout=120,
        )
        if response.status_code == 429 and attempt < 2:
            retry_after = response.headers.get("retry-after")
            reset_tokens = response.headers.get("x-ratelimit-reset-tokens")
            delay = None
            for value in (retry_after, reset_tokens):
                if value:
                    match = re.search(r"([0-9]+(?:\.[0-9]+)?)", value)
                    if match:
                        delay = float(match.group(1))
                        break
            if delay is None:
                match = re.search(r"try again in ([0-9]+(?:\.[0-9]+)?)s", response.text, flags=re.I)
                delay = float(match.group(1)) if match else 15.0
            delay = min(max(delay + 1.0, 1.0), 60.0)
            st.info(f"Limite temporário da Groq. Vou tentar novamente em {delay:.0f} segundos…")
            time.sleep(delay)
            continue
        if response.status_code >= 400:
            raise RuntimeError(f"Groq retornou {response.status_code}: {response.text}")
        content = response.json()["choices"][0]["message"].get("content") or ""
        return content.strip()

    raise RuntimeError("A Groq manteve o limite de tokens ativo após novas tentativas. Aguarde e tente novamente.")


def _call_groq(form_answers, meeting_notes):
    api_key = _secret("GROQ_API_KEY")
    if not api_key:
        raise RuntimeError("GROQ_API_KEY não configurada nos Secrets do Streamlit.")
    if not form_answers.strip():
        raise RuntimeError("Não há respostas preenchidas para gerar o parecer.")

    system = """
Você é assistente técnico do Núcleo de Inovação Tecnológica – NIT/IFSC e deve redigir uma minuta de parecer preliminar de patenteabilidade.

REGRAS:
- As perguntas e respostas preenchidas da planilha abaixo são a fonte principal e obrigatória do parecer. Use efetivamente os dados informados, preservando nomes, percentuais, datas, descrição e referências.
- Não ignore respostas preenchidas relevantes e não substitua a descrição específica por texto genérico.
- Respostas vazias e 'Não aplicável' já foram removidas. Não invente o conteúdo que falta.
- Use as observações da reunião apenas como informação complementar, identificando-as como relato interno.
- Avalie a tecnologia para proteção patentária em caráter preliminar. Use as palavras-chave e as referências de anterioridade registradas na planilha como material informado, sem afirmar que fez uma busca oficial nem que confirmou novidade ou atividade inventiva.
- Diferencie fatos registrados de avaliação técnica preliminar. Se os dados não permitirem uma conclusão, declare a limitação.
- Escreva em português claro, sem criar leis, testes, empresas, concorrentes ou resultados que não apareçam nas fontes.

Estruture o parecer conforme as informações disponíveis:
PARECER TÉCNICO – NIT/IFSC
I – Identificação e inventores
II – Descrição da criação
III – Análise preliminar de proteção e anterioridade
IV – Maturidade e desenvolvimento
V – Aplicações e potencial de exploração
VI – Titularidade e aspectos de propriedade intelectual
VII – Riscos e recomendações
VIII – Conclusão
PARECER DO NIT
"""
    user = f"""
RESPOSTAS PREENCHIDAS DO FORMULÁRIO (coluna, pergunta e resposta):
{form_answers}

OBSERVAÇÕES DA REUNIÃO / COMPLEMENTOS DO NIT:
{meeting_notes.strip() if meeting_notes.strip() else '(Nenhuma informação complementar fornecida.)'}

Redija o parecer usando as respostas específicas acima nas seções correspondentes.
"""
    result = _groq_request(
        api_key,
        [{"role": "system", "content": system}, {"role": "user", "content": user}],
        max_tokens=2600,
    )
    if not result:
        raise RuntimeError("A IA retornou um parecer vazio. Tente novamente.")
    return result


def _add_markdown_like(doc, text):
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            doc.add_paragraph()
            continue
        if line.startswith("### "):
            paragraph = doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            paragraph = doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            paragraph = doc.add_heading(line[2:], level=1)
        elif re.match(r"^[IVX]+\s*[–-]", line):
            paragraph = doc.add_heading(line, level=2)
        else:
            paragraph = doc.add_paragraph()
            parts = re.split(r"(\*\*.*?\*\*)", line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    paragraph.add_run(part[2:-2]).bold = True
                else:
                    paragraph.add_run(part)
        paragraph.paragraph_format.space_after = Pt(6)


def gerar_docx(parecer):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Pt(45)
    section.bottom_margin = Pt(45)
    section.left_margin = Pt(55)
    section.right_margin = Pt(55)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("INSTITUTO FEDERAL DE SANTA CATARINA – IFSC\n")
    run.bold = True
    run.font.size = Pt(13)
    run = title.add_run("NÚCLEO DE INOVAÇÃO TECNOLÓGICA – NIT")
    run.bold = True
    run.font.size = Pt(12)
    doc.add_paragraph()
    _add_markdown_like(doc, parecer)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def _save_history(numero_patente, documento_nome, form_text, notes, parecer):
    try:
        payload = {
            "numero_patente": numero_patente or None,
            "documento_nome": documento_nome or None,
            "formulario_extraido": form_text or None,
            "observacoes_reuniao": notes or None,
            "parecer_gerado": parecer or None,
            "modelo": MODEL,
            "criado_em": datetime.now().isoformat(),
        }
        db._request(
            "POST",
            db._endpoint("pareceres_patenteabilidade"),
            headers=db._headers("return=minimal"),
            json=payload,
        )
        return True
    except Exception:
        return False


def _load_history(limit=100, numero_patente=None):
    params = f"select=*&order=criado_em.desc&limit={int(limit)}"
    if numero_patente and numero_patente.strip():
        from urllib.parse import quote
        params += f"&numero_patente=eq.{quote(numero_patente.strip(), safe='')}"
    data = db._request(
        "GET",
        f"{db._endpoint('pareceres_patenteabilidade')}?{params}",
        headers=db._headers(),
    )
    return data or []


def _delete_history(item_id):
    from urllib.parse import quote
    db._request(
        "DELETE",
        f"{db._endpoint('pareceres_patenteabilidade')}?id=eq.{quote(str(item_id), safe='')}",
        headers=db._headers("return=minimal"),
    )


def _render_history():
    st.divider()
    st.subheader("🗂️ Histórico de pareceres de patenteabilidade")
    st.caption("Pareceres e respostas filtradas da planilha ficam salvos no Supabase.")
    busca = st.text_input(
        "🔎 Buscar por número do processo / PI",
        key="busca_historico_parecer",
        placeholder="Digite o número do processo ou deixe vazio para ver os mais recentes",
    )
    try:
        registros = _load_history(100, busca)
        if not registros:
            st.info("Nenhum parecer registrado para este filtro." if busca.strip() else "Nenhum parecer registrado ainda.")
            return
        for row in registros:
            row_id = row.get("id")
            date = str(row.get("criado_em") or "")[:19].replace("T", " ")
            number = row.get("numero_patente") or "Sem número informado"
            document = row.get("documento_nome") or "Sem planilha"
            with st.expander(f"{date} — {number} — {document}"):
                st.markdown(f"**Processo / PI:** {number}")
                st.markdown(f"**Planilha:** {document}")
                st.markdown(f"**Modelo de IA:** {row.get('modelo') or '-'}")
                st.markdown("**Observações da reunião:**")
                st.write(row.get("observacoes_reuniao") or "Nenhuma observação registrada.")
                st.markdown("**Respostas da planilha usadas:**")
                with st.expander("Ver perguntas e respostas filtradas"):
                    st.text_area(
                        "Respostas",
                        row.get("formulario_extraido") or "",
                        height=280,
                        disabled=True,
                        key=f"planilha_hist_{row_id}",
                    )
                st.markdown("**Parecer gerado:**")
                st.markdown(row.get("parecer_gerado") or "-")
                if st.button("📥 Carregar este parecer", key=f"load_hist_{row_id}"):
                    st.session_state["parecer_patenteabilidade_resultado"] = row.get("parecer_gerado") or ""
                    st.session_state["parecer_patenteabilidade_dados"] = row.get("formulario_extraido") or ""
                    st.session_state["parecer_patenteabilidade_notas_resgatadas"] = row.get("observacoes_reuniao") or ""
                    st.success("Parecer carregado nesta sessão.")
                if st.button("🗑️ Excluir este histórico", key=f"del_hist_{row_id}"):
                    _delete_history(row_id)
                    st.success("Histórico excluído.")
                    st.rerun()
    except Exception as exc:
        st.error(f"Não foi possível consultar o histórico no Supabase: {exc}")


def _render_result():
    parecer = st.session_state.get("parecer_patenteabilidade_resultado")
    if not parecer:
        return
    st.divider()
    st.subheader("📋 Minuta do Parecer Técnico")
    data_used = st.session_state.get("parecer_patenteabilidade_dados")
    if data_used:
        with st.expander("🔗 Conferir perguntas e respostas usadas no parecer"):
            st.text_area("Dados enviados à análise", data_used, height=300, disabled=True, key="parecer_dados_resultado")
    st.markdown(parecer)
    number = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        st.session_state.get("parecer_patenteabilidade_numero", "") or "parecer_patenteabilidade",
    ).strip("_")
    st.download_button(
        "📄 Baixar parecer em Word (.docx)",
        gerar_docx(parecer),
        f"Parecer_Patenteabilidade_NIT_IFSC_{number or 'NIT'}.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True,
    )
    st.download_button(
        "📝 Baixar parecer em texto (.txt)",
        parecer.encode("utf-8"),
        f"Parecer_Patenteabilidade_NIT_IFSC_{number or 'NIT'}.txt",
        "text/plain",
        use_container_width=True,
    )


def render():
    st.title("⚖️ Parecer de Patenteabilidade – NIT/IFSC")
    st.caption("Carregue a planilha do formulário. A linha de perguntas será associada à linha de respostas.")

    arquivo = st.file_uploader(
        "📊 Planilha do Formulário de Notificação de Invenção (Excel ou ODS)",
        type=["xlsx", "xls", "xlsm", "xlsb", "xltx", "xltm", "ods"],
        key="planilha_patenteabilidade",
    )
    numero = st.text_input("Número do processo / PI (opcional)", key="parecer_numero")
    notas = st.text_area(
        "📝 Observações da reunião com o inventor / complementos do NIT",
        height=200,
        value=st.session_state.pop("parecer_patenteabilidade_notas_resgatadas", ""),
        key="parecer_notas_reuniao",
        placeholder="Registre informações complementares obtidas pelo NIT, se houver.",
    )

    if not arquivo:
        _render_history()
        _render_result()
        return

    try:
        all_pairs = _read_spreadsheet(arquivo)
    except Exception as exc:
        st.error(f"Não foi possível ler a planilha: {exc}")
        _render_history()
        return

    patent_form = _is_patent_form(all_pairs)
    if patent_form is not True:
        st.error(
            "Este módulo gera parecer somente quando a resposta ao campo 'Invenção' identifica "
            "um pedido de patente ou modelo de utilidade. Confira o tipo da criação na planilha."
        )
        _render_history()
        _render_result()
        return

    selected_pairs = _filter_patent_pairs(all_pairs)
    if not selected_pairs:
        st.error("Não restaram respostas preenchidas para o tipo de criação selecionado.")
        _render_history()
        return

    form_answers = _format_pairs(selected_pairs)
    st.success(f"{len(selected_pairs)} respostas de patente selecionadas. Colunas EI–UO e datas de envio/última ação foram excluídas.")
    with st.expander("🔎 Conferir respostas que serão enviadas para análise", expanded=True):
        import pandas as pd
        st.dataframe(pd.DataFrame(selected_pairs), use_container_width=True, hide_index=True)

    signature = hashlib.sha256(
        arquivo.getvalue()
        + b"\0" + notas.encode("utf-8")
        + b"\0" + numero.encode("utf-8")
    ).hexdigest()
    if st.session_state.get("parecer_patenteabilidade_signature") != signature:
        st.session_state["parecer_patenteabilidade_signature"] = signature
        st.session_state.pop("parecer_patenteabilidade_resultado", None)
        st.session_state.pop("parecer_patenteabilidade_dados", None)

    if st.button("🤖 Gerar parecer a partir destas respostas", type="primary", use_container_width=True):
        with st.spinner("A IA está analisando as respostas preenchidas do formulário..."):
            try:
                parecer = _call_groq(form_answers, notas)
                st.session_state["parecer_patenteabilidade_resultado"] = parecer
                st.session_state["parecer_patenteabilidade_dados"] = form_answers
                st.session_state["parecer_patenteabilidade_numero"] = numero
                history_saved = _save_history(numero, arquivo.name, form_answers, notas, parecer)
                if not history_saved:
                    st.warning("Parecer gerado, mas não salvo no histórico do Supabase. Verifique a conexão.")
            except Exception as exc:
                st.error(f"Erro ao gerar o parecer: {exc}")
                return

    _render_result()
    _render_history()
